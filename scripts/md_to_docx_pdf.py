#!/usr/bin/env python3
"""Render a Chinese-language markdown research report to DOCX and PDF.

This VM has neither pandoc nor LibreOffice, so DOCX is built with python-docx and
PDF with reportlab. Both need a CJK font with a real bold weight; bootstrap_fonts()
instances one from the Noto Sans SC variable font if it is not already present.

Usage:
    python3 scripts/md_to_docx_pdf.py <report.md> [--outdir DIR] [--title TITLE]
"""

from __future__ import annotations

import argparse
import html
import re
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path

FONT_DIR = Path.home() / ".local/share/fonts"
REGULAR = FONT_DIR / "NotoSansSC-Regular.ttf"
BOLD = FONT_DIR / "NotoSansSC-Bold.ttf"
VF_URL = "https://github.com/notofonts/noto-cjk/raw/main/Sans/Variable/TTF/Subset/NotoSansSC-VF.ttf"
DOCX_FONT = "Noto Sans SC"


def bootstrap_fonts() -> None:
    """Ensure static Regular/Bold CJK TTFs exist (reportlab rejects CFF/OTF and
    variable fonts default to Thin, so the weights must be instanced)."""
    if REGULAR.exists() and BOLD.exists():
        return
    FONT_DIR.mkdir(parents=True, exist_ok=True)
    vf = Path("/tmp/NotoSansSC-VF.ttf")
    if not vf.exists():
        subprocess.run(["curl", "-sSL", "--max-time", "180", "-o", str(vf), VF_URL], check=True)
    from fontTools.ttLib import TTFont as FTFont
    from fontTools.varLib import instancer

    for weight, dest in ((400, REGULAR), (700, BOLD)):
        f = FTFont(str(vf))
        instancer.instantiateVariableFont(f, {"wght": weight}, inplace=True)
        f.save(str(dest))


# --------------------------------------------------------------------------- #
# Markdown parsing
# --------------------------------------------------------------------------- #

@dataclass
class Block:
    kind: str  # h | p | ul | ol | table | quote | hr | code
    level: int = 0
    text: str = ""
    items: list[str] = field(default_factory=list)
    header: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


TABLE_SEP = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")


def split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    cells = re.split(r"(?<!\\)\|", line)
    return [c.strip().replace("\\|", "|") for c in cells]


def parse_md(text: str) -> list[Block]:
    lines = text.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    i = 0
    n = len(lines)

    while i < n:
        raw = lines[i]
        stripped = raw.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            buf: list[str] = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(Block("code", text="\n".join(buf)))
            continue

        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            blocks.append(Block("hr"))
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            blocks.append(Block("h", level=len(m.group(1)), text=m.group(2).strip()))
            i += 1
            continue

        # Table: a pipe row followed by a delimiter row.
        if "|" in stripped and i + 1 < n and TABLE_SEP.match(lines[i + 1]) and "|" in lines[i + 1]:
            header = split_row(stripped)
            i += 2
            rows: list[list[str]] = []
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i]))
                i += 1
            width = max([len(header)] + [len(r) for r in rows]) if rows else len(header)
            header += [""] * (width - len(header))
            rows = [r + [""] * (width - len(r)) for r in rows]
            blocks.append(Block("table", header=header, rows=rows))
            continue

        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(Block("quote", text=" ".join(x for x in buf if x)))
            continue

        bullet = re.match(r"^\s*[-*+]\s+(.*)$", raw)
        number = re.match(r"^\s*\d+[.)]\s+(.*)$", raw)
        if bullet or number:
            kind = "ul" if bullet else "ol"
            pattern = r"^\s*[-*+]\s+(.*)$" if bullet else r"^\s*\d+[.)]\s+(.*)$"
            items: list[str] = []
            while i < n:
                mm = re.match(pattern, lines[i])
                if mm:
                    items.append(mm.group(1).strip())
                    i += 1
                elif lines[i].strip() and lines[i].startswith((" ", "\t")) and items:
                    items[-1] += " " + lines[i].strip()  # continuation line
                    i += 1
                else:
                    break
            blocks.append(Block(kind, items=items))
            continue

        buf = [stripped]
        i += 1
        while i < n and lines[i].strip() and not re.match(
            r"^\s*(#{1,6}\s|[-*+]\s|\d+[.)]\s|>|```|\||-{3,}$)", lines[i]
        ):
            buf.append(lines[i].strip())
            i += 1
        blocks.append(Block("p", text=" ".join(buf)))

    return blocks


# --------------------------------------------------------------------------- #
# Inline markup
# --------------------------------------------------------------------------- #

LINK = re.compile(r"!?\[([^\]]*)\]\(([^)]+)\)")


def inline_segments(text: str) -> list[tuple[str, bool, bool, bool]]:
    """Split inline markdown into (text, bold, italic, code) runs."""
    text = LINK.sub(lambda m: m.group(1) or m.group(2), text)
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.I)
    token = re.compile(r"(\*\*\*|\*\*|__|\*|_|`)")
    out: list[tuple[str, bool, bool, bool]] = []
    bold = italic = code = False
    pos = 0
    while pos < len(text):
        m = token.search(text, pos)
        if not m:
            out.append((text[pos:], bold, italic, code))
            break
        if m.start() > pos:
            out.append((text[pos:m.start()], bold, italic, code))
        tok = m.group(1)
        if code and tok != "`":
            out.append((tok, bold, italic, code))
        elif tok == "`":
            code = not code
        elif tok in ("**", "__"):
            bold = not bold
        elif tok == "***":
            bold = not bold
            italic = not italic
        else:
            # A lone '*' or '_' inside a word (e.g. snake_case) is literal.
            prev = text[m.start() - 1] if m.start() else " "
            nxt = text[m.end()] if m.end() < len(text) else " "
            if prev.isalnum() and nxt.isalnum():
                out.append((tok, bold, italic, code))
            else:
                italic = not italic
        pos = m.end()
    return [(t, b, i, c) for t, b, i, c in out if t]


def to_rl(text: str) -> str:
    """Inline markdown -> reportlab mini-HTML."""
    parts = []
    for seg, bold, italic, code in inline_segments(text):
        s = html.escape(seg).replace("\n", " ")
        if code:
            s = f'<font face="NotoSansSC-Regular" color="#9c2c2c">{s}</font>'
        if italic:
            s = f"<i>{s}</i>"
        if bold:
            s = f"<b>{s}</b>"
        parts.append(s)
    return "".join(parts) or "&nbsp;"


def display_width(text: str) -> int:
    import unicodedata
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #

def build_docx(blocks: list[Block], title: str, dest: Path) -> None:
    import docx
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = docx.Document()

    st = doc.styles["Normal"]
    st.font.name = DOCX_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)

    section = doc.sections[0]
    section.left_margin = section.right_margin = docx.shared.Cm(2.0)

    def style_run(run, *, bold=False, italic=False, mono=False, size=None, color=None):
        name = "Consolas" if mono else DOCX_FONT
        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        run.font.bold = bold
        run.font.italic = italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)

    def emit_inline(par, text, *, size=None, bold_all=False, color=None):
        for seg, bold, italic, code in inline_segments(text):
            run = par.add_run(seg)
            style_run(
                run,
                bold=bold or bold_all,
                italic=italic,
                mono=code,
                size=size,
                color=(0x9C, 0x2C, 0x2C) if code else color,
            )

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    emit_inline(head, title, size=20, bold_all=True)

    heading_sizes = {1: 17, 2: 14.5, 3: 12.5, 4: 11.5, 5: 11, 6: 10.5}

    for b in blocks:
        if b.kind == "h":
            if b.level == 1 and b.text.strip() == title.strip():
                continue
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(10 if b.level <= 2 else 6)
            par.paragraph_format.space_after = Pt(4)
            emit_inline(par, b.text, size=heading_sizes.get(b.level, 11), bold_all=True,
                        color=(0x1F, 0x3B, 0x63) if b.level <= 2 else None)
        elif b.kind == "p":
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(4)
            emit_inline(par, b.text)
        elif b.kind in ("ul", "ol"):
            for idx, item in enumerate(b.items, 1):
                clean = re.sub(r"^\[[ xX]\]\s*", "", item)
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = docx.shared.Cm(0.75)
                par.paragraph_format.space_after = Pt(2)
                prefix = f"{idx}. " if b.kind == "ol" else "• "
                style_run(par.add_run(prefix))
                emit_inline(par, clean)
        elif b.kind == "quote":
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = docx.shared.Cm(0.75)
            emit_inline(par, b.text, color=(0x55, 0x55, 0x55))
        elif b.kind == "hr":
            doc.add_paragraph()
        elif b.kind == "code":
            par = doc.add_paragraph()
            style_run(par.add_run(b.text), mono=True, size=9)
        elif b.kind == "table":
            cols = len(b.header)
            table = doc.add_table(rows=1 + len(b.rows), cols=cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, cell_text in enumerate(b.header):
                cell = table.cell(0, c)
                cell.text = ""
                emit_inline(cell.paragraphs[0], cell_text, size=9.5, bold_all=True)
            for r, row in enumerate(b.rows, 1):
                for c, cell_text in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = ""
                    emit_inline(cell.paragraphs[0], cell_text, size=9.5)
            doc.add_paragraph()

    doc.save(str(dest))


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #

def build_pdf(blocks: list[Block], title: str, dest: Path) -> int:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    pdfmetrics.registerFont(TTFont("NotoSansSC-Regular", str(REGULAR)))
    pdfmetrics.registerFont(TTFont("NotoSansSC-Bold", str(BOLD)))
    pdfmetrics.registerFontFamily(
        "NotoSansSC-Regular", normal="NotoSansSC-Regular", bold="NotoSansSC-Bold",
        italic="NotoSansSC-Regular", boldItalic="NotoSansSC-Bold",
    )

    base = ParagraphStyle(
        "body", fontName="NotoSansSC-Regular", fontSize=9.5, leading=15,
        alignment=TA_LEFT, spaceAfter=4, wordWrap="CJK",
    )
    title_style = ParagraphStyle("title", parent=base, fontName="NotoSansSC-Bold",
                                 fontSize=19, leading=26, alignment=TA_CENTER, spaceAfter=14)
    heads = {
        1: ParagraphStyle("h1", parent=base, fontName="NotoSansSC-Bold", fontSize=15.5,
                          leading=22, spaceBefore=12, spaceAfter=6, textColor=colors.HexColor("#1f3b63")),
        2: ParagraphStyle("h2", parent=base, fontName="NotoSansSC-Bold", fontSize=13,
                          leading=19, spaceBefore=10, spaceAfter=5, textColor=colors.HexColor("#1f3b63")),
        3: ParagraphStyle("h3", parent=base, fontName="NotoSansSC-Bold", fontSize=11.5,
                          leading=17, spaceBefore=8, spaceAfter=4),
        4: ParagraphStyle("h4", parent=base, fontName="NotoSansSC-Bold", fontSize=10.5,
                          leading=16, spaceBefore=6, spaceAfter=3),
    }
    heads[5] = heads[6] = heads[4]
    list_style = ParagraphStyle("li", parent=base, leftIndent=12, bulletIndent=2, spaceAfter=2)
    quote_style = ParagraphStyle("quote", parent=base, leftIndent=12,
                                 textColor=colors.HexColor("#555555"), borderPadding=2)
    code_style = ParagraphStyle("code", parent=base, fontSize=8.5, leading=12,
                                backColor=colors.HexColor("#f4f4f4"), borderPadding=4)
    th_style = ParagraphStyle("th", parent=base, fontName="NotoSansSC-Bold",
                              fontSize=8, leading=11, spaceAfter=0, textColor=colors.white)
    td_style = ParagraphStyle("td", parent=base, fontSize=8, leading=11, spaceAfter=0)

    doc = SimpleDocTemplate(
        str(dest), pagesize=A4,
        leftMargin=16 * mm, rightMargin=16 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
        title=title, author="Opus 深度研究",
    )
    avail = doc.width

    story: list = [Paragraph(to_rl(title), title_style),
                   HRFlowable(width="100%", color=colors.HexColor("#1f3b63"), thickness=1.1),
                   Spacer(1, 8)]

    for b in blocks:
        if b.kind == "h":
            if b.level == 1 and b.text.strip() == title.strip():
                continue
            story.append(Paragraph(to_rl(b.text), heads.get(b.level, heads[4])))
        elif b.kind == "p":
            story.append(Paragraph(to_rl(b.text), base))
        elif b.kind in ("ul", "ol"):
            for idx, item in enumerate(b.items, 1):
                clean = re.sub(r"^\[([ xX])\]\s*", lambda m: "☑ " if m.group(1) in "xX" else "☐ ", item)
                bullet = f"{idx}." if b.kind == "ol" else "•"
                story.append(Paragraph(to_rl(clean), list_style, bulletText=bullet))
        elif b.kind == "quote":
            story.append(Paragraph(to_rl(b.text), quote_style))
        elif b.kind == "hr":
            story.append(Spacer(1, 4))
            story.append(HRFlowable(width="100%", color=colors.HexColor("#cccccc"), thickness=0.6))
            story.append(Spacer(1, 4))
        elif b.kind == "code":
            story.append(Paragraph(to_rl(b.text).replace(" ", "&nbsp;"), code_style))
        elif b.kind == "table":
            cols = len(b.header)
            if cols == 0:
                continue
            # Share width by longest cell, but damped and floored so a verbose column
            # cannot squeeze a short one down to one character per line.
            weights = []
            for c in range(cols):
                cells = [b.header[c]] + [r[c] for r in b.rows]
                longest = max((display_width(x) for x in cells), default=1)
                weights.append(max(6.0, min(float(longest), 44.0)) ** 0.72)
            total = sum(weights) or 1.0
            widths = [avail * w / total for w in weights]

            data = [[Paragraph(to_rl(h), th_style) for h in b.header]]
            for row in b.rows:
                data.append([Paragraph(to_rl(cell), td_style) for cell in row])

            tbl = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3b63")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b0b8c4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
            ]))
            story.append(Spacer(1, 3))
            story.append(tbl)
            story.append(Spacer(1, 7))

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("NotoSansSC-Regular", 8)
        canvas.setFillColor(colors.HexColor("#888888"))
        canvas.drawCentredString(A4[0] / 2.0, 9 * mm, f"{title}  ·  第 {canvas.getPageNumber()} 页")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)

    import pymupdf
    with pymupdf.open(str(dest)) as d:
        return d.page_count


# --------------------------------------------------------------------------- #

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("markdown")
    ap.add_argument("--outdir")
    ap.add_argument("--title")
    args = ap.parse_args()

    src = Path(args.markdown)
    text = src.read_text(encoding="utf-8")
    outdir = Path(args.outdir) if args.outdir else src.parent
    outdir.mkdir(parents=True, exist_ok=True)

    bootstrap_fonts()
    blocks = parse_md(text)

    title = args.title
    if not title:
        first_h1 = next((b.text for b in blocks if b.kind == "h" and b.level == 1), None)
        title = first_h1 or src.stem
    title = re.sub(r"[*`]", "", title).strip()

    docx_path = outdir / f"{src.stem}.docx"
    pdf_path = outdir / f"{src.stem}.pdf"
    build_docx(blocks, title, docx_path)
    pages = build_pdf(blocks, title, pdf_path)

    tables = sum(1 for b in blocks if b.kind == "table")
    print(
        f"{src.name}: {len(text)} chars, {len(blocks)} blocks, {tables} tables\n"
        f"  -> {docx_path}  ({docx_path.stat().st_size:,} bytes)\n"
        f"  -> {pdf_path}  ({pdf_path.stat().st_size:,} bytes, {pages} pages)"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
